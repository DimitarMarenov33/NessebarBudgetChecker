"""Tests for the budget-execution parsers, using real sample files in
`data/samples/` as fixtures:

- `nesebar_budget_execution_aug2026.xlsx` / `nesebar_cash_execution_B1_2026_8.xls`:
  2026, EUR-denominated (the forms the parsers were originally built on).
- `nesebar_budget_execution_feb2021_BGN.xlsx`: 2021 capital ledger, BGN,
  no explicit currency marker (period-based fallback).
- `nesebar_cash_execution_B1_2020_1_BGN.xls`: 2020 monthly B1, BGN, with
  its own "(в лева)" marker.
- `nesebar_cash_execution_B3_2019_1_BGN.xls`: 2019 Q1 quarterly B3 (same
  `OTCHET` schema as B1, cumulative Jan-Mar), BGN.

Conversion rate: 1 EUR = 1.95583 BGN (Bulgaria's official euro-adoption
rate, 2026-01-01).
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
import xlrd
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

import nessebar_budget.db.budget_models  # noqa: F401 -- registers the line-item tables
from nessebar_budget.db.budget_models import BudgetLineItem, CashExecutionLine
from nessebar_budget.db.budget_repo import (
    ReportCandidate,
    group_candidates,
    parse_budget_reports,
    preference_key,
    upsert_budget_report,
)
from nessebar_budget.db.models import Base
from nessebar_budget.parsers import budget_b1, budget_capital
from nessebar_budget.parsers.budget_b1 import BGN_PER_EUR as B1_BGN_PER_EUR
from nessebar_budget.parsers.budget_b1 import parse_budget_b1
from nessebar_budget.parsers.budget_capital import BGN_PER_EUR as CAPITAL_BGN_PER_EUR
from nessebar_budget.parsers.budget_capital import parse_budget_capital
from nessebar_budget.scrapers.nesebar_site import classify_filename

SAMPLES = Path(__file__).resolve().parent.parent / "data" / "samples"
CAPITAL_XLSX = SAMPLES / "nesebar_budget_execution_aug2026.xlsx"
CAPITAL_XLSX_2021_BGN = SAMPLES / "nesebar_budget_execution_feb2021_BGN.xlsx"
B1_XLS = SAMPLES / "nesebar_cash_execution_B1_2026_8.xls"
B1_XLS_2020_BGN = SAMPLES / "nesebar_cash_execution_B1_2020_1_BGN.xls"
B3_XLS_2019_BGN = SAMPLES / "nesebar_cash_execution_B3_2019_1_BGN.xls"

RATE = 1.95583


def eur(bgn: float) -> float:
    return round(bgn / RATE, 2)


# -- fixtures -----------------------------------------------------------------


@pytest.fixture(scope="module")
def capital_result():
    return parse_budget_capital(CAPITAL_XLSX)


@pytest.fixture(scope="module")
def capital_result_2021_bgn():
    return parse_budget_capital(CAPITAL_XLSX_2021_BGN, period="2021-02")


@pytest.fixture(scope="module")
def b1_result():
    return parse_budget_b1(B1_XLS)


@pytest.fixture(scope="module")
def b1_result_2020_bgn():
    return parse_budget_b1(B1_XLS_2020_BGN)


@pytest.fixture(scope="module")
def b3_result_2019_bgn():
    return parse_budget_b1(B3_XLS_2019_BGN)


# -- capital ledger: 2026 EUR sample (original layout) ------------------------


def test_capital_ledger_grand_total(capital_result):
    """The "Общо" sheet's ОБЩО row should carry the municipality-wide total
    Сметна стойност (estimated_total) of 24,990,398 EUR."""
    rows = capital_result["rows"]
    totals = [
        r
        for r in rows
        if r["unit"] == "Общо" and r["object_code"] == "ОБЩО" and r["extra_json"]["row_type"] == "grand_total"
    ]
    assert len(totals) == 1
    assert totals[0]["estimated_total"] == pytest.approx(24990398)
    assert totals[0]["currency"] == "EUR"


def test_capital_ledger_captures_named_object(capital_result):
    """A specific object row ("Леки автомобили за Общинска администрация")
    with a planned/estimated cost of 100,000 EUR should be present."""
    rows = capital_result["rows"]
    matches = [
        r
        for r in rows
        if r["object_name"] and "Леки автомобили за Общинска администрация" in r["object_name"]
    ]
    assert len(matches) == 1
    row = matches[0]
    assert row["unit"] == "Общо"
    assert row["estimated_total"] == pytest.approx(100000)
    assert row["object_code"] == "1122"


def test_capital_ledger_no_validation_mismatches(capital_result):
    """Function subtotals should reconcile with the sum of their objects
    (the sample file is well-formed)."""
    assert capital_result["mismatches"] == []


def test_capital_ledger_parses_every_sheet(capital_result):
    units = {r["unit"] for r in capital_result["rows"]}
    assert "Общо" in units
    assert len(units) > 1  # кметства/schools sheets also parsed


def test_capital_ledger_not_unsupported(capital_result):
    assert capital_result["unsupported"] is None


def test_capital_ledger_eur_marker_no_original_block(capital_result):
    """2026 sample has its own "Сумите са в EUR!" marker: no conversion
    happens, and no `extra_json["original"]` block is recorded."""
    totals = [
        r
        for r in capital_result["rows"]
        if r["unit"] == "Общо" and r["object_code"] == "ОБЩО"
    ]
    extra = totals[0]["extra_json"]
    assert extra["original_currency"] == "EUR"
    assert "original" not in extra
    assert "conversion_rate" not in extra


# -- capital ledger: 2021 BGN sample (currency fallback + header detection) --


def test_capital_ledger_2021_bgn_currency_fallback(capital_result_2021_bgn):
    """The 2021 workbook has no explicit currency marker; the period
    (2021-02, < 2026-01) should drive the BGN fallback."""
    assert capital_result_2021_bgn["unsupported"] is None
    totals = [
        r
        for r in capital_result_2021_bgn["rows"]
        if r["unit"] == "Общо" and r["object_code"] == "ОБЩО"
    ]
    assert len(totals) == 1
    extra = totals[0]["extra_json"]
    assert extra["original_currency"] == "BGN"
    assert extra["conversion_rate"] == pytest.approx(CAPITAL_BGN_PER_EUR)


def test_capital_ledger_2021_bgn_conversion_known_cell(capital_result_2021_bgn):
    """The "Общо" grand total's Сметна стойност is 39,380,680 BGN on the raw
    sheet; the parsed row should carry that converted to EUR (2 dp), with
    the original BGN figure preserved in extra_json."""
    totals = [
        r
        for r in capital_result_2021_bgn["rows"]
        if r["unit"] == "Общо" and r["object_code"] == "ОБЩО"
    ]
    row = totals[0]
    assert row["currency"] == "EUR"
    assert row["extra_json"]["original_currency"] == "BGN"
    assert row["extra_json"]["conversion_rate"] == pytest.approx(1.95583)
    assert row["estimated_total"] == pytest.approx(eur(39380680))


def test_capital_ledger_2021_header_detected_by_label(capital_result_2021_bgn):
    """Sanity check that the label-based header locator (not a fixed row
    number) found real data: paragraph subtotal §51-00 should be present
    with a nonzero estimated_total."""
    rows = capital_result_2021_bgn["rows"]
    para = [
        r
        for r in rows
        if r["unit"] == "Общо" and r["paragraph"] == "5100" and r["extra_json"]["row_type"] == "paragraph_subtotal"
    ]
    assert len(para) == 1
    assert para[0]["estimated_total"] and para[0]["estimated_total"] > 0


# -- capital ledger: synthetic shifted-column workbook (layout variant) ------


def _write_shifted_capital_workbook(path: Path) -> None:
    """Build a minimal capital-ledger workbook whose core columns are
    shifted one column to the right of the 2026/2021 layout, and whose
    currency marker says "лева" instead of "EUR" -- a stand-in for an
    unseen older-year layout variant, to prove the header/column locator
    works off label text rather than fixed positions.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Общо"

    ws.cell(row=2, column=2, value="ОБЩИНА Несебър")
    ws.cell(row=3, column=3, value="план/отчет за периода:  2022 Март")
    ws.cell(row=4, column=2, value="Сумите са в лева")

    # Header row, shifted +1 column vs. the real-world layout.
    header_row = 6
    ws.cell(row=header_row, column=2, value="§")
    ws.cell(row=header_row, column=3, value="Информация за наименованието на обекта")
    ws.cell(row=header_row, column=4, value="Година начало - година край")
    ws.cell(row=header_row, column=5, value="Сметна стойност")
    ws.cell(row=header_row, column=6, value="Усвоено до края на предходната година")
    ws.cell(row=header_row, column=7, value="Уточнен план")
    ws.cell(row=header_row, column=8, value="Усвоено към отчетния период")

    ws.cell(row=header_row + 2, column=3, value="ОБЩО")
    ws.cell(row=header_row + 2, column=5, value=195583)  # = 100,000 EUR
    ws.cell(row=header_row + 2, column=6, value=0)
    ws.cell(row=header_row + 2, column=7, value=0)
    ws.cell(row=header_row + 2, column=8, value=0)

    ws.cell(row=header_row + 3, column=2, value="5100")
    ws.cell(row=header_row + 3, column=3, value="Основен ремонт на дълготрайни материални активи")
    ws.cell(row=header_row + 3, column=5, value=195583)
    ws.cell(row=header_row + 3, column=6, value=0)
    ws.cell(row=header_row + 3, column=7, value=0)
    ws.cell(row=header_row + 3, column=8, value=0)

    wb.save(path)


def test_capital_ledger_header_detection_handles_shifted_columns(tmp_path):
    path = tmp_path / "shifted_capital.xlsx"
    _write_shifted_capital_workbook(path)

    result = parse_budget_capital(path, period="2022-03")
    assert result["unsupported"] is None

    totals = [r for r in result["rows"] if r["object_code"] == "ОБЩО"]
    assert len(totals) == 1
    assert totals[0]["estimated_total"] == pytest.approx(100000)
    assert totals[0]["extra_json"]["original_currency"] == "BGN"

    para = [r for r in result["rows"] if r["paragraph"] == "5100"]
    assert len(para) == 1
    assert para[0]["estimated_total"] == pytest.approx(100000)


def test_capital_ledger_unsupported_when_no_header_label_found(tmp_path):
    """A workbook with no "Сметна стойност" label anywhere (e.g. a misfiled
    report of a different kind) should be classified as unsupported, not
    silently mis-parsed."""
    path = tmp_path / "not_a_capital_ledger.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="this is some other kind of report")
    wb.save(path)

    result = parse_budget_capital(path, period="2022-03")
    assert result["rows"] == []
    assert result["unsupported"] is not None


# -- B1: 2026 EUR sample (original layout) ------------------------------------


def test_b1_expenditure_table_nonempty_and_numeric(b1_result):
    """The B1 parser should return a non-empty expenditure-by-paragraph
    table with numeric plan/actual values."""
    rows = b1_result["rows"]
    expenditure = [r for r in rows if r["section"] == "разходи"]
    assert len(expenditure) > 0

    numeric_rows = [r for r in expenditure if isinstance(r["actual_ytd"], float)]
    assert len(numeric_rows) > 0
    assert any(r["actual_ytd"] and r["actual_ytd"] > 0 for r in numeric_rows)

    salaries = [r for r in expenditure if r["paragraph"] == "100"]
    assert len(salaries) == 1
    assert salaries[0]["actual_ytd"] > 0


def test_b1_revenue_table_nonempty(b1_result):
    revenue = [r for r in b1_result["rows"] if r["section"] == "приходи"]
    assert len(revenue) > 0


def test_b1_documents_skipped_pages(b1_result):
    """The sheet has many more report pages than revenue+expenditure; the
    parser should say so rather than silently dropping them."""
    assert b1_result["skipped"]


def test_b1_eur_marker_no_original_block(b1_result):
    expenditure = [r for r in b1_result["rows"] if r["section"] == "разходи" and r["paragraph"] == "100"]
    extra = expenditure[0]["extra_json"]
    assert extra["original_currency"] == "EUR"
    assert "original" not in extra


# -- B1: 2020 BGN sample (currency marker + conversion) -----------------------


def test_b1_2020_bgn_marker_detected(b1_result_2020_bgn):
    expenditure = [r for r in b1_result_2020_bgn["rows"] if r["section"] == "разходи" and r["paragraph"] == "100"]
    assert len(expenditure) == 1
    extra = expenditure[0]["extra_json"]
    assert extra["original_currency"] == "BGN"
    assert extra["conversion_rate"] == pytest.approx(B1_BGN_PER_EUR)


def test_b1_2020_bgn_conversion_known_cell(b1_result_2020_bgn):
    """§100 (salaries) actual_total on the raw BGN sheet should convert to
    EUR (2 dp), with the raw BGN figure preserved in extra_json."""
    expenditure = [r for r in b1_result_2020_bgn["rows"] if r["section"] == "разходи" and r["paragraph"] == "100"]
    row = expenditure[0]
    assert row["extra_json"]["original_currency"] == "BGN"
    assert row["extra_json"]["conversion_rate"] == pytest.approx(1.95583)
    assert row["actual_ytd"] and row["actual_ytd"] > 0
    assert row["currency"] == "EUR"


def test_b1_2020_period_is_pre2026(b1_result_2020_bgn):
    assert b1_result_2020_bgn["rows"][0]["period"] == "2020-01"


# -- B3: quarterly report accepted by the same B1 parser ---------------------


def test_b3_quarterly_parses_same_schema(b3_result_2019_bgn):
    """A B3 (quarterly) workbook shares B1's exact `OTCHET` schema and
    should be parsed by the same function with no special-casing."""
    rows = b3_result_2019_bgn["rows"]
    assert rows
    expenditure = [r for r in rows if r["section"] == "разходи" and r["paragraph"] == "100"]
    assert len(expenditure) == 1
    assert expenditure[0]["actual_ytd"] and expenditure[0]["actual_ytd"] > 0


def test_b3_quarterly_period_uses_quarter_end_not_year_start(b3_result_2019_bgn):
    """`_infer_period` must pick the quarter-end ("до") date, not the
    year-start ("от") date that a naive first-date-found scan would pick
    (Q1 2019's "от" is 2019-01-01, which would misreport as January)."""
    assert b3_result_2019_bgn["rows"][0]["period"] == "2019-03"


def test_b3_quarterly_bgn_currency_detected(b3_result_2019_bgn):
    expenditure = [r for r in b3_result_2019_bgn["rows"] if r["section"] == "разходи" and r["paragraph"] == "100"]
    assert expenditure[0]["extra_json"]["original_currency"] == "BGN"


def test_classify_filename_b3_is_its_own_kind():
    """B3/IB3 (quarterly) files share B1/IB1's schema (same parser) but are
    their own kinds, so a B3 and the same quarter-end month's B1 can be told
    apart and de-duplicated. The B3's third number is the quarter, so the
    filename month is the quarter-end month."""
    info = classify_filename("B3_2019_1_5206.xls")
    assert info["kind"] == "B3"
    assert (info["year"], info["month"], info["quarter"]) == (2019, 3, 1)

    info = classify_filename("B3_2018_4_5206.xls")
    assert (info["kind"], info["year"], info["month"]) == ("B3", 2018, 12)

    info = classify_filename("IB3_2019_1_5206_K33.xls")
    assert info["kind"] == "IB3_K33"

    assert classify_filename("B1_2019_1_5206.xls")["kind"] == "B1"
    assert classify_filename("IB1_2019_1_5206_DES.xls")["kind"] == "IB1_DES"


@pytest.mark.parametrize(
    "filename",
    [
        "Budget_2018_5206.xlsx",
        "Budget_2020_5206.xlsx",
        "NaturiPokazateli_2019_5206.xlsx",
        "kp01.01.16.xlsx",
    ],
)
def test_classify_filename_non_ledger_xlsx_is_other(filename):
    assert classify_filename(filename)["kind"] == "other"


def test_non_ledger_xlsx_still_downloaded():
    """Reclassifying them as "other" must not stop them being fetched."""
    from nessebar_budget.scrapers.nesebar_site import _should_download

    assert _should_download({"kind": "other", "filename": "Budget_2018_5206.xlsx"})
    assert _should_download({"kind": "B3", "filename": "B3_2019_1_5206.xls"})
    assert _should_download({"kind": "IB3_RA", "filename": "IB3_2019_1_5206_RA.xls"})
    assert not _should_download({"kind": "other", "filename": "regUV2019.pdf"})


# -- filename classification (unchanged behaviour) ----------------------------


@pytest.mark.parametrize(
    ("filename", "expected_kind", "expected_year", "expected_month"),
    [
        ("B1_2026_8_5206.xls", "B1", 2026, 8),
        ("IB1_2026_8_5206_DES.xls", "IB1_DES", 2026, 8),
        ("IB1_2026_8_5206_K33.xls", "IB1_K33", 2026, 8),
        ("august2026.xlsx", "capital_xlsx", 2026, 8),
        ("june26.xlsx", "capital_xlsx", 2026, 6),
    ],
)
def test_classify_filename(filename, expected_kind, expected_year, expected_month):
    info = classify_filename(filename)
    assert info["kind"] == expected_kind
    assert info["year"] == expected_year
    assert info["month"] == expected_month


def test_classify_filename_unknown_falls_back_to_other():
    info = classify_filename("regUV2019.pdf")
    assert info["kind"] == "other"


# -- period from the file's own header (not the nesebar.bg heading month) -----


@pytest.mark.parametrize(
    ("path", "expected"),
    [
        (CAPITAL_XLSX, "2026-08"),
        (CAPITAL_XLSX_2021_BGN, "2021-02"),
    ],
)
def test_capital_detect_period(path, expected):
    assert budget_capital.detect_period(path) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("план/отчет за периода:  2022 Юни", "2022-06"),
        ("ПЛАН/ОТЧЕТ ЗА ПЕРИОДА:2023   септември", "2023-09"),
        ("план/отчет\nза периода: 2024 МАЙ", "2024-05"),
        ("за периода: Декември 2021", "2021-12"),
        ("план/отчет за периода:", None),
        ("2022 Юнии", None),
    ],
)
def test_capital_period_text_variants(text, expected):
    assert budget_capital._period_from_text(text) == expected


def test_capital_detect_period_none_for_non_ledger(tmp_path):
    path = tmp_path / "Budget_2018_5206.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "BUDGET"
    wb.active["A1"] = "ОТЧЕТНИ ДАННИ"
    wb.save(path)
    assert budget_capital.detect_period(path) is None


@pytest.mark.parametrize(
    ("path", "expected"),
    [
        # Q1 2019: "от" 2019-01-01, "до" 2019-03-31 -> the quarter-end month.
        (B3_XLS_2019_BGN, "2019-03"),
        (B1_XLS_2020_BGN, "2020-01"),
        (B1_XLS, "2026-08"),
    ],
)
def test_cash_detect_period_uses_do_date(path, expected):
    assert budget_b1.detect_period(path) == expected


# -- provenance: every parsed row points at its sheet and Excel row ----------


@pytest.mark.parametrize("path", [CAPITAL_XLSX, CAPITAL_XLSX_2021_BGN])
def test_capital_rows_carry_source_sheet_and_row(path, capital_result, capital_result_2021_bgn):
    result = capital_result if path == CAPITAL_XLSX else capital_result_2021_bgn
    rows = result["rows"]
    assert rows
    assert all("source_sheet" in r["extra_json"] and "source_row" in r["extra_json"] for r in rows)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        obj = next(r for r in rows if r["extra_json"]["row_type"] == "object")
        ws = wb[obj["extra_json"]["source_sheet"]]
        cells = next(
            ws.iter_rows(
                min_row=obj["extra_json"]["source_row"],
                max_row=obj["extra_json"]["source_row"],
                max_col=2,
                values_only=True,
            )
        )
        assert str(cells[0]).strip() == obj["object_code"]
        assert str(cells[1]).strip() == obj["object_name"]
    finally:
        wb.close()


@pytest.mark.parametrize("path", [B1_XLS, B1_XLS_2020_BGN, B3_XLS_2019_BGN])
def test_cash_rows_carry_source_sheet_and_row(path):
    rows = parse_budget_b1(path)["rows"]
    assert rows
    assert all(r["extra_json"]["source_sheet"] == "OTCHET" for r in rows)
    assert all(isinstance(r["extra_json"]["source_row"], int) for r in rows)
    sheet = xlrd.open_workbook(path).sheet_by_name("OTCHET")
    for row in rows[:20]:
        excel_row = row["extra_json"]["source_row"]
        names = {str(sheet.cell_value(excel_row - 1, col)).strip() for col in (2, 3)}
        assert row["name"] in names


# -- one file per (family, period): preference order --------------------------


def _cand(report_id, kind, name, period="2022-06", mtime=0.0):
    return ReportCandidate(
        report_id=report_id, kind=kind, file_path=f"/cache/2022/07/{name}", period=period, mtime=mtime
    )


def test_capital_preference_monthly_over_quarterly_over_copies_over_mtime():
    monthly_copy = _cand(1, "capital_xlsx", "Месечен отчет за 2022 Юни 5206 Несебър(1).xlsx")
    quarterly = _cand(2, "capital_xlsx", "Тримесечен отчет за 2022 Юни 5206 Несебър.xlsx")
    quarterly_latin = _cand(3, "capital_xlsx", "3mesechen-mart.xlsx")
    neutral = _cand(4, "capital_xlsx", "42022.xlsx")
    otchet = _cand(5, "capital_xlsx", "otchet2022dec5206.xlsx")
    m_otchet_old = _cand(6, "capital_xlsx", "m-otchet-aug.xlsx", mtime=100.0)
    m_otchet_new = _cand(7, "capital_xlsx", "m-otchet-5206.xlsx", mtime=200.0)

    ranked = sorted(
        [quarterly, quarterly_latin, neutral, monthly_copy, otchet, m_otchet_old, m_otchet_new],
        key=preference_key,
    )
    ids = [c.report_id for c in ranked]
    # monthly-named, original before "(1)", newest mtime first
    assert ids[:4] == [7, 6, 5, 1]
    # then unmarked, then quarterly-named ("Тримесечен" must not count as monthly)
    assert ids[4] == 4
    assert set(ids[5:]) == {2, 3}


def test_capital_preference_original_over_reupload_copy():
    original = _cand(1, "capital_xlsx", "Месечен отчет за 2023 Март 5206 Несебър.xlsx", mtime=1.0)
    copy = _cand(2, "capital_xlsx", "Месечен отчет за 2023 Март 5206 Несебър(1).xlsx", mtime=9.0)
    spaced = _cand(3, "capital_xlsx", "Месечен отчет за 2022 Август 5206 Несебър (1).xlsx")
    assert min([copy, spaced, original], key=preference_key) is original


def test_cash_preference_b3_over_b1_and_grouping():
    b1 = _cand(10, "B1", "B1_2022_6_5206.xls", period="2022-06", mtime=50.0)
    b3 = _cand(11, "B3", "B3_2022_2_5206.xls", period="2022-06", mtime=1.0)
    other_month = _cand(12, "B1", "B1_2022_7_5206.xls", period="2022-07")
    capital = _cand(13, "capital_xlsx", "Месечен отчет за 2022 Юни 5206 Несебър.xlsx")
    no_period = _cand(14, "B1", "B1_x.xls", period=None)
    not_parseable = _cand(15, "IB3_DES", "IB3_2022_2_5206_DES.xls")

    groups = group_candidates([b1, b3, other_month, capital, no_period, not_parseable])
    assert [c.report_id for c in groups[("cash", "2022-06")]] == [11, 10]
    assert [c.report_id for c in groups[("cash", "2022-07")]] == [12]
    assert [c.report_id for c in groups[("capital", "2022-06")]] == [13]
    assert [c.report_id for c in groups[("cash", "?14")]] == [14]
    assert all(15 not in [c.report_id for c in g] for g in groups.values())


# -- parse_budget_reports: periods, dedupe, cleanup, provenance (temp DB) -----


@pytest.fixture
def budget_db(tmp_path):
    """Temp DB with cached-report rows as the old scraper left them: periods
    = nesebar.bg heading month, B3 files catalogued as kind "B1"."""
    import shutil

    cache = tmp_path / "cache"
    cache.mkdir()

    def copy(src: Path, name: str) -> str:
        dest = cache / name
        shutil.copy(src, dest)
        return str(dest)

    non_ledger = cache / "Budget_2018_5206.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "BUDGET"
    wb.active["A1"] = "ОТЧЕТНИ ДАННИ ПО ЕБК"
    wb.save(non_ledger)

    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    session = Session(engine)
    specs = [
        # (kind, heading period, file)
        ("B1", "2019-04", copy(B3_XLS_2019_BGN, "B3_2019_1_5206.xls")),  # legacy kind
        ("B1", "2019-03", copy(B3_XLS_2019_BGN, "B1_2019_3_5206.xls")),  # same period
        ("B1", "2020-01", copy(B1_XLS_2020_BGN, "B1_2020_1_5206.xls")),
        ("IB1_K33", "2019-04", copy(B3_XLS_2019_BGN, "IB3_2019_1_5206_K33.xls")),
        ("capital_xlsx", "2021-03", copy(CAPITAL_XLSX_2021_BGN, "Отчет за 2021 Февруари.xlsx")),
        ("capital_xlsx", "2021-02", copy(CAPITAL_XLSX_2021_BGN, "Отчет за 2021 Февруари(1).xlsx")),
        ("capital_xlsx", "2024-02", str(non_ledger)),
    ]
    reports = []
    for kind, period, path in specs:
        report = upsert_budget_report(
            session,
            {
                "url": f"https://www.nesebar.bg/03-2019/{Path(path).name}",
                "kind": kind,
                "period": period,
                "file_path": path,
                "parsed_json": {"period_source": "heading"},
            },
        )
        reports.append(report)
    # A stale row left by the old heading-based parse of the duplicate B1.
    session.add(CashExecutionLine(report_id=reports[1].id, period="2019-03", section="разходи"))
    session.commit()
    yield session, reports
    session.close()


def test_parse_budget_reports_periods_dedupe_and_cleanup(budget_db):
    session, reports = budget_db
    b3, b1_dup, b1_2020, ib3, cap, cap_copy, non_ledger = reports
    urls = {r.id: (r.url, r.file_path) for r in reports}

    summary = parse_budget_reports(session)
    session.commit()

    # Kinds: B3/IB3 refined from the legacy B1/IB1 kinds.
    assert b3.kind == "B3"
    assert ib3.kind == "IB3_K33"

    # Period from content; heading guess kept.
    assert b3.period == "2019-03"
    assert b3.parsed_json["period_source"] == "content"
    assert b3.parsed_json["heading_period"] == "2019-04"
    assert cap.period == "2021-02"
    assert cap.parsed_json["heading_period"] == "2021-03"

    # One file per family and period: B3 kept over B1, original over "(1)".
    assert b1_dup.parsed_json["skipped"]["duplicate_of"] == b3.id
    assert cap_copy.parsed_json["skipped"]["duplicate_of"] == cap.id
    assert summary.skipped_duplicates == 2

    # Non-ledger xlsx: parse attempted, reclassified, reason kept.
    assert non_ledger.kind == "other"
    assert "Сметна стойност" in non_ledger.parsed_json["other_reason"]

    # Rows only for the chosen files, at the detected period; stale row gone.
    cash = session.scalars(select(CashExecutionLine)).all()
    assert {r.report_id for r in cash} == {b3.id, b1_2020.id}
    assert {r.period for r in cash if r.report_id == b3.id} == {"2019-03"}
    capital = session.scalars(select(BudgetLineItem)).all()
    assert {r.report_id for r in capital} == {cap.id}
    assert {r.period for r in capital} == {"2021-02"}

    # Provenance on every stored row; url/file_path untouched.
    assert all(r.extra_json["source_sheet"] and r.extra_json["source_row"] for r in cash + capital)
    assert {r.id: (r.url, r.file_path) for r in reports} == urls

    # Idempotent: a second run changes nothing; --rebuild gives the same rows.
    counts = (len(cash), len(capital))
    again = parse_budget_reports(session)
    session.commit()
    assert again.rows_deleted == 0 and again.periods_changed == 0
    rebuilt = parse_budget_reports(session, rebuild=True)
    session.commit()
    assert rebuilt.files_parsed == summary.files_parsed == 3
    assert (
        len(session.scalars(select(CashExecutionLine)).all()),
        len(session.scalars(select(BudgetLineItem)).all()),
    ) == counts


def test_rescrape_keeps_content_period(budget_db):
    """A later scrape upsert (heading period, filename kind) must not move a
    file back to its heading month or undo a content-based "other"."""
    session, reports = budget_db
    parse_budget_reports(session)
    b3, non_ledger = reports[0], reports[6]
    upsert_budget_report(
        session,
        {"url": b3.url, "kind": "B3", "period": "2019-04", "parsed_json": {"heading": "x"}},
    )
    upsert_budget_report(
        session,
        {"url": non_ledger.url, "kind": "capital_xlsx", "period": "2024-02", "parsed_json": {}},
    )
    assert (b3.period, b3.kind, b3.parsed_json["heading_period"]) == ("2019-03", "B3", "2019-04")
    assert b3.parsed_json["row_count"] > 0
    assert non_ledger.kind == "other"
